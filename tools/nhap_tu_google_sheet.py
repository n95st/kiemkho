#!/usr/bin/env python3
"""One-off import of the legacy Google Sheets (downloaded as .xlsx) into the Supabase schema `kk`.

  python3 tools/nhap_tu_google_sheet.py --trung-tam trungtam.xlsx --dot DOT01=dot1.xlsx [--dot DOT02=dot2.xlsx ...] > nhap.sql
  psql "$DB_URL" -v ON_ERROR_STOP=1 -f nhap.sql

Prints one transaction (begin; ... commit;) on stdout; warnings go to stderr.
The SQL refuses to run when kk.dong or kk.tai_khoan already has rows.
"""
import argparse
import datetime as dt
import re
import sys

import openpyxl

KHO_OK = ("CHAN", "LE")
ROLES = ("DEM", "TO_TRUONG", "IT", "TAP")
HOC_ST = ("CHUA_NHAN", "DANG_DEM", "XONG")
CHUNK = 500
# Sheet columns (0-based) as in legacy/Code.gs: C = DONG columns.
C = dict(ID=0, TG=1, KHO=2, MA=3, TEN=4, LO=5, HSD=6, HOC=7, HOP=8, QC=9, LE=10, SL=11, DVT=12, NGUOI=13, TK=14,
         TT=15, ST=16, THAY_BOI=17, GHI=18, CANH_BAO=19, REQ=20, CAP_NHAT=21, SE_THAY=22)


def warn(sheet, row, msg):
    print("CẢNH BÁO %s dòng %s: %s" % (sheet, row, msg), file=sys.stderr)


# ---------------------------------------------------------------- cell helpers
def txt(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y %H:%M") if isinstance(v, dt.datetime) else v.strftime("%d/%m/%Y")
    return str(v).replace("\x00", "").strip()


def lit(s):
    return "'" + str(s).replace("'", "''") + "'"


def sq(v):
    """SQL literal for text; empty -> null."""
    t = txt(v)
    return lit(t) if t else "null"


def num(v, sheet, row, what):
    """Number literal via repr; empty or unparseable -> null."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return "null"
    if isinstance(v, bool):
        return "null"
    if isinstance(v, (int, float)):
        return repr(int(v)) if v == int(v) else repr(float(v))
    try:
        f = float(str(v).strip().replace(",", "."))
    except ValueError:
        warn(sheet, row, "%s không phải số: %r" % (what, v))
        return "null"
    return repr(int(f)) if f == int(f) else repr(f)


def parse_hsd(s):
    """Same rules as kk.parse_hsd: d/m/yyyy, d-m-yy, m/yyyy (-> last day)."""
    s = s.strip()
    m = re.fullmatch(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)
    try:
        if m:
            y = int(m[3]) + (2000 if int(m[3]) < 100 else 0)
            return dt.date(y, int(m[2]), int(m[1])) if 2000 <= y <= 2100 else None
        m = re.fullmatch(r"(\d{1,2})[/\-.](\d{4})", s)
        if m:
            y, mo = int(m[2]), int(m[1])
            first = dt.date(y, mo, 1)
            return (first.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)
    except ValueError:
        return None
    return None


def date_lit(v, sheet, row):
    if v is None or txt(v) == "":
        return "null"
    d = v.date() if isinstance(v, dt.datetime) else v if isinstance(v, dt.date) else parse_hsd(str(v))
    if d is None:
        warn(sheet, row, "HSD không đọc được, để trống: %r" % (v,))
        return "null"
    return lit(d.isoformat())


def ts_lit(v, sheet, row):
    """Local Vietnam time -> timestamptz literal; None if unreadable."""
    if isinstance(v, dt.datetime):
        d = v
    elif isinstance(v, dt.date):
        d = dt.datetime(v.year, v.month, v.day)
    else:
        s = txt(v)
        m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})(?:\s+(\d{1,2}):(\d{2})(?::(\d{2}))?)?", s)
        if not m:
            if s:
                warn(sheet, row, "thời điểm không đọc được: %r" % (v,))
            return None
        try:
            d = dt.datetime(int(m[3]), int(m[2]), int(m[1]), int(m[4] or 0), int(m[5] or 0), int(m[6] or 0))
        except ValueError:
            warn(sheet, row, "thời điểm không đọc được: %r" % (v,))
            return None
    return lit(d.strftime("%Y-%m-%d %H:%M:%S") + "+07")


def norm(v):
    """Same rules as kk.lo_text / kk.hoc_key: upper, trim, collapse spaces."""
    return re.sub(r"\s+", " ", txt(v).upper()).strip()


def so_of(v, prefix):
    m = re.fullmatch(prefix + r"(\d+)", txt(v))
    return int(m[1]) if m else None


def ints(v):
    return [int(x) for x in re.findall(r"\d+", txt(v))]


def rows_of(wb, name, width, first=2):
    """Rows (padded to width) of a sheet, skipping the header; [] when the sheet is missing."""
    if name not in wb.sheetnames:
        return
    for i, r in enumerate(wb[name].iter_rows(min_row=first, values_only=True), start=first):
        r = tuple(r)[:width]
        yield i, r + (None,) * (width - len(r))


def insert(table, cols, rows, tail=""):
    for i in range(0, len(rows), CHUNK):
        part = rows[i:i + CHUNK]
        yield "insert into kk.%s (%s) values\n  %s%s;" % (table, ", ".join(cols), ",\n  ".join("(" + ", ".join(r) + ")" for r in part), tail)


def cfg_of(wb):
    return {txt(r[0]): r[1] for _, r in rows_of(wb, "CAU_HINH", 2, first=1) if txt(r[0])}


# ---------------------------------------------------------------- sheets
def hub_sql(wb, out):
    users, seen_tk, seen_pin = [], set(), set()
    for i, r in rows_of(wb, "NGUOI_DEM", 6):
        tk, pin = txt(r[0]), txt(r[2])
        if not tk or not pin:
            continue
        if isinstance(r[2], (int, float)) and len(pin) < 6:
            warn("NGUOI_DEM", i, "PIN %s là số, có thể đã mất số 0 đầu – kiểm lại với người dùng" % tk)
        if not re.fullmatch(r"\d{4,8}", pin) or tk in seen_tk or pin in seen_pin:
            warn("NGUOI_DEM", i, "bỏ qua tài khoản %s (PIN sai dạng hoặc trùng tài khoản/PIN)" % tk)
            continue
        seen_tk.add(tk), seen_pin.add(pin)
        role = txt(r[3]) if txt(r[3]) in ROLES else "DEM"
        active = r[4] is True or txt(r[4]).upper() == "TRUE"
        users.append([lit(tk), lit(txt(r[1]) or tk), lit(pin), lit(role), "true" if active else "false", lit(txt(r[5]))])
    out += insert("tai_khoan", ["tk", "ten", "pin", "vai_tro", "hoat_dong", "ghi_chu"], users)

    dm = {}
    for i, r in rows_of(wb, "DANH_MUC", 8):
        ma = txt(r[0])
        if not ma:
            continue
        kho = txt(r[5]).upper()
        if kho and kho not in KHO_OK:
            warn("DANH_MUC", i, "kho mặc định %r lạ, để trống" % kho)
            kho = ""
        gtin = list(dict.fromkeys(t for t in re.split(r"[\s,;]+", txt(r[6])) if t))
        dm[ma] = [lit(ma), lit(txt(r[1])), lit(txt(r[2])), num(r[3], "DANH_MUC", i, "Quy cách"), lit(txt(r[4])),
                  lit(kho) if kho else "null",
                  "array[%s]::text[]" % ", ".join(lit(g) for g in gtin) if gtin else "'{}'::text[]", lit(txt(r[7]))]
    out += insert("thuoc", ["ma", "ten", "dvt_nhap", "quy_cach", "dvt", "kho_mac_dinh", "gtin", "nguon"], list(dm.values()))

    cfg = cfg_of(wb)
    for k in ("LO_MAC_DINH", "HSD_MAC_DINH"):
        v = txt(cfg.get(k)).lstrip("'")
        if v:
            out.append("insert into kk.cau_hinh (khoa, gia_tri) values (%s, %s::jsonb) on conflict (khoa) do update set gia_tri = excluded.gia_tri;"
                       % (lit(k), lit('"' + v.replace("\\", "\\\\").replace('"', '\\"') + '"')))

    dots = []
    for i, r in rows_of(wb, "DOT", 10):
        did = txt(r[0])
        if not did:
            continue
        kho = [k for k in re.split(r"[\s,;]+", txt(r[3]).upper()) if k in KHO_OK] or list(KHO_OK)
        dots.append(dict(id=did, ten=txt(r[1]) or did, st="DONG" if txt(r[2]).upper() == "DONG" else "MO", kho=kho,
                         ngay=ts_lit(r[4], "DOT", i), nguoi=txt(r[5]), so_voi=txt(r[6]), ghi=txt(r[9])))
    return dots


def dot_sql(wb, dot, out):
    name = dot["id"]
    cfg = cfg_of(wb)
    if not dot["so_voi"] and txt(cfg.get("SO_VOI")) not in ("", "(không so sánh)"):
        dot["so_voi"] = txt(cfg.get("SO_VOI"))

    dong, seen, max_so, seen_req = [], set(), 0, set()
    for i, r in rows_of(wb, "DONG", 23):
        if not txt(r[C["ID"]]):
            continue
        so = so_of(r[C["ID"]], "D")
        sl = num(r[C["SL"]], "DONG", i, "SL")
        if so is None or so in seen or sl == "null" or not txt(r[C["MA"]]):
            warn("DONG", i, "bỏ qua dòng %r (mã dòng sai/trùng, thiếu mã thuốc hoặc SL)" % txt(r[C["ID"]]))
            continue
        seen.add(so)
        max_so = max(max_so, so)
        tg = ts_lit(r[C["TG"]], "DONG", i)
        if tg is None:
            warn("DONG", i, "thiếu thời điểm, dùng giờ nhập")
        req = txt(r[C["REQ"]])
        if req in seen_req:
            warn("DONG", i, "mã yêu cầu %s trùng, để trống" % req)
            req = ""
        seen_req.add(req)
        se_thay = ints(r[C["SE_THAY"]])
        thay_boi = ints(r[C["THAY_BOI"]])
        dong.append([lit(name), str(so), "false", tg or "now()", lit(txt(r[C["KHO"]]).upper()), lit(txt(r[C["MA"]])),
                     lit(txt(r[C["TEN"]])), lit(norm(r[C["LO"]])), date_lit(r[C["HSD"]], "DONG", i), lit(norm(r[C["HOC"]])),
                     num(r[C["HOP"]], "DONG", i, "Số hộp"), num(r[C["QC"]], "DONG", i, "1 hộp có mấy"),
                     num(r[C["LE"]], "DONG", i, "Số lẻ"), sl, lit(txt(r[C["DVT"]])), lit(txt(r[C["NGUOI"]])),
                     lit(txt(r[C["TK"]])), lit(txt(r[C["TT"]])), lit(txt(r[C["ST"]])),
                     str(thay_boi[0]) if thay_boi else "null", lit(txt(r[C["GHI"]])), lit(txt(r[C["CANH_BAO"]])),
                     lit(req) if req else "null", lit(txt(r[C["CAP_NHAT"]])),
                     "array[%s]::int[]" % ", ".join(map(str, se_thay)) if se_thay else "'{}'::int[]"])

    moi, seen_m, max_m = [], set(), 0
    for i, r in rows_of(wb, "THUOC_MOI", 12):
        so = so_of(r[0], "M")
        if not txt(r[0]):
            continue
        if so is None or so in seen_m:
            warn("THUOC_MOI", i, "bỏ qua dòng %r (mã sai/trùng)" % txt(r[0]))
            continue
        seen_m.add(so)
        max_m = max(max_m, so)
        moi.append([lit(name), str(so), ts_lit(r[1], "THUOC_MOI", i) or "now()", lit(txt(r[2])), lit(txt(r[3])), lit(txt(r[4])),
                    num(r[5], "THUOC_MOI", i, "Quy cách"), lit(txt(r[6])), lit(txt(r[7]).upper()), lit(txt(r[8])),
                    lit(txt(r[9]) or "CHO_DUYET"), lit(txt(r[10])), lit(txt(r[11]))])

    hoc, seen_h = [], set()
    for i, r in rows_of(wb, "HOC", 8):
        kho, h = txt(r[0]).upper(), norm(r[1])
        if not kho or not h or (kho, h) in seen_h:
            continue
        seen_h.add((kho, h))
        hoc.append([lit(name), "false", lit(kho), lit(h), lit(txt(r[2]) if txt(r[2]) in HOC_ST else "CHUA_NHAN"), lit(txt(r[3])),
                    ts_lit(r[4], "HOC", i) or "null", ts_lit(r[5], "HOC", i) or "null", lit(txt(r[6])), lit(txt(r[7]))])

    ss = []
    for i, r in rows_of(wb, "LAN1", 9):
        if not txt(r[1]):
            continue
        ss.append([lit(name), lit(txt(r[0]).upper()), lit(txt(r[1])), lit(txt(r[2])), lit(norm(r[3])), date_lit(r[4], "LAN1", i),
                   num(r[5], "LAN1", i, "SL"), lit(txt(r[6])), lit(norm(r[7])), lit(txt(r[8]))])

    out.append("insert into kk.dot (id, ten, trang_thai, kho, ngay_tao, nguoi_tao, so_voi, ghi_chu, so_dong, so_thuoc_moi) values (%s);" % ", ".join([
        lit(name), lit(dot["ten"]), lit(dot["st"]), "array[%s]::text[]" % ", ".join(lit(k) for k in dot["kho"]), dot["ngay"] or "now()",
        lit(dot["nguoi"]), lit(dot["so_voi"]), lit(dot["ghi"]), str(max_so), str(max_m)]))
    out += insert("dot_kho", ["dot_id", "kho", "trang_thai", "cap_nhat"], [
        [lit(name), lit(k), lit("DA_CHOT" if txt(cfg.get("TRANG_THAI_" + k)) == "DA_CHOT" else "DANG_KIEM"), lit(txt(cfg.get("TRANG_THAI_%s_LUC" % k)))]
        for k in dot["kho"]])
    out += insert("dong", ["dot_id", "so", "tap", "tg", "kho", "ma", "ten", "lo", "hsd", "hoc", "hop", "qc", "le", "sl", "dvt", "nguoi", "tk",
                           "thao_tac", "trang_thai", "thay_boi", "ghi", "canh_bao", "req_id", "cap_nhat", "se_thay"], dong)
    out += insert("thuoc_moi", ["dot_id", "so", "tg", "ma_de_xuat", "ten", "dvt_nhap", "qc_goc", "dvt", "kho", "nguoi", "trang_thai",
                                "ma_chinh_thuc", "ghi_chu"], moi)
    out += insert("hoc", ["dot_id", "tap", "kho", "hoc", "trang_thai", "nguoi", "bat_dau", "xong_luc", "nguoi_xong", "ghi"], hoc)
    out += insert("so_sanh", ["dot_id", "kho", "ma", "ten", "lo", "hsd", "sl", "dvt", "hoc", "nguoi"], ss)


GUARD = ("do $$ begin if exists (select 1 from kk.dong) or exists (select 1 from kk.tai_khoan) then "
         "raise exception 'Cơ sở dữ liệu đã có dữ liệu – dừng nhập'; end if; end $$;")
WIPE = ("-- --xoa-du-lieu-thu: wipe seed/test data first (local testing only)\n"
        "delete from kk.so_sanh; delete from kk.hoc; delete from kk.thuoc_moi; delete from kk.dong; delete from kk.dot_kho;\n"
        "delete from kk.dot; delete from kk.thuoc; delete from kk.tai_khoan; delete from kk.nhat_ky; delete from kk.pin_sai;")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--trung-tam", required=True, help="Trung tâm workbook (NGUOI_DEM, DANH_MUC, DOT, CAU_HINH)")
    ap.add_argument("--dot", action="append", default=[], metavar="MÃ=FILE", help="đợt workbook, e.g. DOT01=dot1.xlsx (repeatable)")
    ap.add_argument("--xoa-du-lieu-thu", action="store_true",
                    help="LOCAL TESTING ONLY: delete the seed data instead of refusing to run on a non-empty database")
    a = ap.parse_args()
    files = {}
    for s in a.dot:
        k, _, f = s.partition("=")
        if not k or not f:
            ap.error("--dot cần dạng MÃ=FILE, nhận %r" % s)
        files[k.strip()] = f

    load = lambda p: openpyxl.load_workbook(p, data_only=True, read_only=True)
    out = ["begin;", WIPE if a.xoa_du_lieu_thu else GUARD,
           "insert into kk.kho (ma, ten, ten_ngan, thu_tu) values ('CHAN', 'Kho chẵn', 'CHẴN', 1), ('LE', 'Kho lẻ', 'LẺ', 2) on conflict do nothing;"]
    dots = hub_sql(load(a.trung_tam), out)
    unknown = set(files) - {d["id"] for d in dots}
    if unknown:
        ap.error("đợt %s không có trong sheet DOT của file Trung tâm" % ", ".join(sorted(unknown)))
    for d in dots:
        if d["id"] not in files:
            warn("DOT", d["id"], "không có file đợt, chỉ nhập dòng đợt (không có dữ liệu đếm)")
            out.append("insert into kk.dot (id, ten, trang_thai, kho, ngay_tao, nguoi_tao, so_voi, ghi_chu) values (%s);" % ", ".join([
                lit(d["id"]), lit(d["ten"]), lit(d["st"]), "array[%s]::text[]" % ", ".join(lit(k) for k in d["kho"]), d["ngay"] or "now()",
                lit(d["nguoi"]), lit(d["so_voi"]), lit(d["ghi"])]))
            out += insert("dot_kho", ["dot_id", "kho"], [[lit(d["id"]), lit(k)] for k in d["kho"]])
        else:
            dot_sql(load(files[d["id"]]), d, out)
    out.append("commit;")
    sys.stdout.reconfigure(encoding="utf-8")
    print("\n".join(out))


if __name__ == "__main__":
    main()
